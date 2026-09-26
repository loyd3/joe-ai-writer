"""从上传文件字节中提取纯文本（txt / md / docx / pdf）。"""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Tuple


_MAX_CHARS = 500_000  # 单文件提取上限，防止极端大文档撑爆内存


def _normalize_text(text: str) -> str:
    text = (text or "").replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # 压缩过多空行
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _from_plain(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gb18030", "gbk", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _from_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise ValueError("服务器未安装 python-docx，无法解析 Word") from e
    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    # 表格单元格也纳入
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ValueError("服务器未安装 pypdf，无法解析 PDF") from e
    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        t = t.strip()
        if t:
            parts.append(t)
    text = "\n\n".join(parts)
    if not text.strip():
        raise ValueError("PDF 未能提取到文字（可能是扫描件，需 OCR）")
    return text


def extract_text_from_bytes(filename: str, data: bytes) -> Tuple[str, str]:
    """
    返回 (cleaned_text, format_label)。
    format_label: txt | md | docx | pdf
    """
    if not data:
        raise ValueError("文件为空")

    name = (filename or "file").strip()
    ext = Path(name).suffix.lower()

    if ext in (".docx",):
        raw = _from_docx(data)
        fmt = "docx"
    elif ext in (".pdf",):
        raw = _from_pdf(data)
        fmt = "pdf"
    elif ext in (".doc",):
        raise ValueError("暂不支持旧版 .doc，请另存为 .docx 或 .pdf")
    elif ext in (
        ".txt",
        ".md",
        ".markdown",
        ".text",
        ".json",
        ".csv",
        ".html",
        ".htm",
        ".log",
        "",
    ):
        raw = _from_plain(data)
        fmt = "md" if ext in (".md", ".markdown") else "txt"
    else:
        # 尝试当纯文本；失败则明确报错
        try:
            raw = _from_plain(data)
            fmt = "txt"
            if "\x00" in raw[:2000]:
                raise ValueError("binary")
        except Exception as e:
            raise ValueError(
                f"不支持的文件类型「{ext or '未知'}」，请使用 txt/md/docx/pdf"
            ) from e

    text = _normalize_text(raw)
    if len(text) > _MAX_CHARS:
        text = text[:_MAX_CHARS]
    if len(text) < 20:
        raise ValueError("提取到的文字过少，请确认文件含可读正文")
    return text, fmt
